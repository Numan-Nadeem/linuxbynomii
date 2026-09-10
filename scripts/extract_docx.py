#!/usr/bin/env python3
"""Extract text + image positions from the Linux course docx."""
import json
import re

from docx import Document
from docx.oxml.ns import qn

SRC = "/home/z/my-project/upload/Linux System Administration Course.docx"
OUT = "/home/z/my-project/extracted/content.json"

doc = Document(SRC)

# Map relationship ids -> image filenames
rels = doc.part.rels
rel_map = {}
for rid, rel in rels.items():
    if "image" in rel.reltype:
        rel_map[rid] = rel.target_ref.split("/")[-1]

def get_images_from_run(run):
    """Return list of image filenames referenced inside a run."""
    imgs = []
    # blip elements
    for blip in run._element.findall(".//" + qn("a:blip")):
        embed = blip.get(qn("r:embed"))
        if embed and embed in rel_map:
            imgs.append(rel_map[embed])
    return imgs

items = []  # list of {type, text/style/images}
for para in doc.paragraphs:
    text = para.text.strip()
    imgs = []
    for run in para.runs:
        imgs.extend(get_images_from_run(run))
    style = para.style.name if para.style else "Normal"
    if imgs:
        items.append({"type": "image", "images": imgs, "style": style})
    if text:
        items.append({"type": "text", "style": style, "text": text})

# Tables
tables = []
for t_idx, table in enumerate(doc.tables):
    rows = []
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        rows.append(cells)
    tables.append({"index": t_idx, "rows": rows})

result = {"paragraphs": items, "tables": tables}

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, indent=1)

# Print summary
print(f"Paragraph items: {len(items)}")
print(f"Tables: {len(tables)}")
print("=" * 60)
for it in items:
    if it["type"] == "image":
        print(f"[IMAGE] {it['images']}")
    else:
        style = it["style"]
        text = it["text"][:110]
        print(f"({style}) {text}")
